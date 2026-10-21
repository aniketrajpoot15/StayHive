import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Brand Colors
    PRIMARY = RGBColor(254, 66, 77)      # StayHive Coral / Red
    DARK_BG = RGBColor(15, 23, 42)       # Slate 900
    DARK_CARD = RGBColor(30, 41, 59)     # Slate 800
    WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    TEXT_DARK = RGBColor(30, 41, 59)
    LIGHT_BG = RGBColor(248, 250, 252)   # Slate 50
    CARD_BG = RGBColor(255, 255, 255)
    ACCENT_BLUE = RGBColor(14, 165, 233)
    BORDER_COLOR = RGBColor(226, 232, 240)

    def add_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="STAYHIVE PROJECT PRESENTATION"):
        # Category / Subtitle tag
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.35))
        p_cat = tb_cat.text_frame.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY
        
        # Main Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = DARK_BG

    def create_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE (Dark Premium Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, DARK_BG)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()

    # Title & Subtitle Box
    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.5), Inches(3.0))
    tf = tbox.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "🏡 StayHive"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = PRIMARY

    p2 = tf.add_paragraph()
    p2.text = "Full-Stack Vacation Rental & Accommodation Platform"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "Inspired by Airbnb | Built with Node.js, Express, MongoDB Atlas, Mapbox & Cloudinary"
    p3.font.size = Pt(16)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(12)

    # Metadata Card at Bottom
    c_meta = create_card(s1, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.2), DARK_CARD, None)
    tb_meta = s1.shapes.add_textbox(Inches(1.1), Inches(5.5), Inches(11.1), Inches(1.0))
    tf_m = tb_meta.text_frame
    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "Presenter: Aniket Rajpoot  |  Project: Full-Stack Web Development"
    p_m1.font.size = Pt(15)
    p_m1.font.bold = True
    p_m1.font.color.rgb = WHITE

    p_m2 = tf_m.add_paragraph()
    p_m2.text = "🌐 Live Demo: https://stay-hive-three.vercel.app/listings  |  GitHub: github.com/aniketrajpoot15/StayHive"
    p_m2.font.size = Pt(13)
    p_m2.font.color.rgb = ACCENT_BLUE
    p_m2.space_before = Pt(6)

    s1.notes_slide.notes_text_frame.text = (
        "Good morning everyone! Today I am presenting StayHive, a comprehensive full-stack vacation rental "
        "and accommodation web platform inspired by Airbnb. It empowers property hosts to list unique stays "
        "and enables travelers to discover places worldwide with interactive maps and authentic reviews. "
        "The project is currently live on Vercel at stay-hive-three.vercel.app."
    )

    # ==========================================
    # SLIDE 2: PROJECT OVERVIEW & MOTIVATION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2, LIGHT_BG)
    add_header(s2, "Executive Overview & Problem Statement")

    # 3 Cards: Problem, Solution, Target Audience
    cols = [
        ("❌ The Problem", 
         ["High barrier for local hosts to showcase distinctive stays.",
          "Lack of visual location context (where is the stay located?).",
          "Unreliable photo hosting & slow page load times.",
          "Fragmented review and rating verification."],
         PRIMARY),
        ("💡 The StayHive Solution", 
         ["Effortless listing creation with real-time forward geocoding.",
          "Interactive Mapbox GL maps embedded with custom location pins.",
          "Optimized cloud media streaming powered by Cloudinary CDN.",
          "Secure user authentication and transparent star reviews."],
         ACCENT_BLUE),
        ("👥 Key Stakeholders", 
         ["Guests / Travelers: Browse, filter by category, explore map coordinates, and leave feedback.",
          "Property Hosts: Add new properties, update pricing/photos, and manage listings.",
          "Admins: Ensure platform integrity, data validation, and clean content."],
         DARK_BG)
    ]

    for i, (col_title, items, border_c) in enumerate(cols):
        left = Inches(0.8 + i * 4.0)
        c = create_card(s2, left, Inches(1.8), Inches(3.7), Inches(5.0), WHITE, border_c)
        tb = s2.shapes.add_textbox(left + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = col_title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = border_c
        
        for item in items:
            pi = tf.add_paragraph()
            pi.text = "• " + item
            pi.font.size = Pt(13)
            pi.font.color.rgb = TEXT_DARK
            pi.space_before = Pt(12)

    s2.notes_slide.notes_text_frame.text = (
        "In this slide, we cover why StayHive was built. Existing platforms are often overly complex or "
        "closed off. StayHive solves this by giving users a clean, responsive Airbnb-like UI where listings "
        "have verified coordinates, dynamic cloud images, and secure owner management."
    )

    # ==========================================
    # SLIDE 3: ARCHITECTURE & SYSTEM DESIGN (MVC)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3, LIGHT_BG)
    add_header(s3, "System Architecture & MVC Design Pattern")

    layers = [
        ("🎨 Presentation Layer (View)", 
         "EJS & EJS-Mate Templates",
         ["Boilerplate layouts & reusable partials (Navbar, Footer, Flash alerts)",
          "Bootstrap 5.3 + Vanilla CSS3 for Airbnb-inspired aesthetics",
          "Mapbox GL JS client-side rendering & Starability animated star reviews",
          "Client-side form validation using HTML5 and Bootstrap classes"]),
        ("⚙️ Business Logic (Controller & Router)", 
         "Node.js & Express.js REST APIs",
         ["RESTful routing pattern for Listings, Reviews, and Users",
          "Express middleware pipeline for auth protection & schema validation",
          "Mapbox SDK forward geocoding integration (Address -> Lat/Long)",
          "Multer & Cloudinary streaming storage engine for image uploads"]),
        ("🗄️ Persistence Layer (Model)", 
         "MongoDB Atlas & Mongoose ODM",
         ["Normalized Mongoose schemas with ObjectId cross-references",
          "Listing Schema with embedded GeoJSON Point geometry",
          "Passport-Local-Mongoose plugin for secure user password hashing",
          "Mongoose post-delete middleware hooks for cascade review cleanup"])
    ]

    for i, (title, subtitle, bullets) in enumerate(layers):
        top = Inches(1.8 + i * 1.7)
        c = create_card(s3, Inches(0.8), top, Inches(11.7), Inches(1.5), WHITE, BORDER_COLOR)
        
        # Tag on left
        tag = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(0.15), Inches(1.5))
        tag.fill.solid()
        tag.fill.fore_color.rgb = PRIMARY if i == 0 else (ACCENT_BLUE if i == 1 else DARK_BG)
        tag.line.fill.background()

        tb = s3.shapes.add_textbox(Inches(1.2), top + Inches(0.15), Inches(11.0), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title + "  —  " + subtitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = DARK_BG

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "✔  " + b
            pb.font.size = Pt(12)
            pb.font.color.rgb = TEXT_DARK
            pb.space_before = Pt(3)

    s3.notes_slide.notes_text_frame.text = (
        "StayHive is architected strictly following the Model-View-Controller (MVC) paradigm. "
        "This ensures clean separation of concerns: EJS templates handle rendering, Express controllers "
        "manage the business rules, and Mongoose schemas protect data integrity."
    )

    # ==========================================
    # SLIDE 4: TECHNOLOGY STACK BREAKDOWN
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4, LIGHT_BG)
    add_header(s4, "Technology Stack & Cloud Ecosystem")

    tech_cards = [
        ("Frontend Technologies",
         [("EJS & EJS-Mate", "Dynamic server-side templating with modular layouts"),
          ("Bootstrap 5.3", "Responsive grid system, modals, forms & dropdowns"),
          ("Vanilla CSS3", "Custom glassmorphism, responsive filters & typography"),
          ("Starability CSS", "Accessible, interactive star rating widgets")]),
        ("Backend & Runtime",
         [("Node.js (v20+)", "Asynchronous, event-driven JavaScript runtime"),
          ("Express.js (v5)", "Fast, robust RESTful routing and middleware stack"),
          ("Joi Validation", "Server-side schema validation for request payloads"),
          ("Passport.js", "Authentication middleware with PBKDF2 salt hashing")]),
        ("Database & Storage",
         [("MongoDB Atlas", "Scalable cloud NoSQL database with high availability"),
          ("Mongoose ODM", "Schema definitions, type casting, validation & hooks"),
          ("Connect-Mongo", "Production session persistence across serverless invocations"),
          ("Cloudinary SDK", "Cloud-based asset storage, transformations & CDN delivery")]),
        ("APIs & Deployment",
         [("Mapbox GL JS", "Interactive web mapping and vector tile rendering"),
          ("Mapbox Geocoding", "Automated forward geocoding from text to coordinates"),
          ("Vercel Serverless", "Zero-config AWS Lambda-powered global edge deployment"),
          ("Git & GitHub", "Source code version control and CI/CD automated deployment")])
    ]

    for idx, (group_title, items) in enumerate(tech_cards):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 2.65)
        
        c = create_card(s4, left, top, Inches(5.7), Inches(2.45), WHITE, BORDER_COLOR)
        
        tb = s4.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), Inches(5.1), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = group_title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = PRIMARY if col == 0 else ACCENT_BLUE

        for name, desc in items:
            pi = tf.add_paragraph()
            pi.text = f"• {name}: {desc}"
            pi.font.size = Pt(12)
            pi.font.color.rgb = TEXT_DARK
            pi.space_before = Pt(4)

    s4.notes_slide.notes_text_frame.text = (
        "Here is the breakdown of the technologies used. Notice the blend of battle-tested open source "
        "libraries with modern cloud services like MongoDB Atlas, Cloudinary, and Mapbox."
    )

    # ==========================================
    # SLIDE 5: CORE FEATURES & CAPABILITIES
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5, LIGHT_BG)
    add_header(s5, "Core Features & Functional Modules")

    features = [
        ("🏠 Full CRUD for Listings",
         "Hosts can create, read, update, and delete properties. Every listing features high-res photos, formatted pricing in INR, detailed descriptions, and city/country locations."),
        ("🗺️ Automated Geocoding & Maps",
         "When a host types an address, Mapbox forward-geocodes it into exact [longitude, latitude] coordinates, plotting an interactive zoomable pin on the listing detail page."),
        ("☁️ Cloud Multi-part Image Uploads",
         "Powered by Multer and Cloudinary. Images uploaded by users are stored securely in the cloud, generating secure CDN URLs for lightning-fast global delivery."),
        ("⭐ Review System & Star Ratings",
         "Verified users can post reviews with 1-to-5 star ratings and feedback. Review authors retain permission to delete their own reviews."),
        ("🏷️ Category Filtering UI",
         "An Airbnb-style horizontal filter bar supporting categories like Trending, Rooms, Iconic Cities, Mountains, Castles, Pools, Camping, Beachfront, and Arctic."),
        ("⚡ Live Tax Toggle Switch",
         "Interactive client-side switch that instantly toggles display prices between base rates and all-inclusive total rates including 18% GST.")
    ]

    for idx, (title, desc) in enumerate(features):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 1.75)

        c = create_card(s5, left, top, Inches(5.7), Inches(1.55), WHITE, BORDER_COLOR)
        tb = s5.shapes.add_textbox(left + Inches(0.25), top + Inches(0.15), Inches(5.2), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = DARK_BG

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(4)

    s5.notes_slide.notes_text_frame.text = (
        "These six core modules form the backbone of StayHive. From geocoding every listing on the fly "
        "to interactive review systems and real-time tax toggles, the user experience closely mirrors Airbnb."
    )

    # ==========================================
    # SLIDE 6: SECURITY, AUTHENTICATION & AUTHORIZATION
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6, LIGHT_BG)
    add_header(s6, "Security Architecture, Authentication & Permissions")

    sec_cards = [
        ("🔐 Authentication (Passport.js)",
         ["Local Strategy integrated with Passport-Local-Mongoose.",
          "PBKDF2 cryptographic hashing with unique random salt per user.",
          "Password strings are never stored in plaintext.",
          "Flash notifications for successful logins, logouts, and invalid credentials."]),
        ("🛡️ Granular Authorization (RBAC)",
         ["isLoggedIn Middleware: Guards private routes (/new, /edit, reviews).",
          "isOwner Middleware: Ensures only the creator can edit or delete a listing.",
          "isReviewAuthor Middleware: Prevents users from deleting reviews posted by others.",
          "UI elements (Edit/Delete buttons) conditionally render based on currUser."]),
        ("🔍 Robust Input Validation (Joi)",
         ["Server-side schema validation defends against malicious payloads or Postman injection.",
          "ListingSchema validates required fields, numeric price bounds, and coordinates.",
          "ReviewSchema enforces 1-5 rating range and non-empty comment strings.",
          "Bootstrap custom form validation delivers immediate visual client feedback."]),
        ("🍪 Session Security & Data Integrity",
         ["connect-mongo session store isolates session data in MongoDB Atlas.",
          "HTTP-Only cookie flags mitigate Cross-Site Scripting (XSS) token theft.",
          "Mongoose Cascade Delete Hook removes all associated reviews when a listing is deleted.",
          "Custom ExpressError class normalizes error status codes and friendly error pages."])
    ]

    for idx, (title, items) in enumerate(sec_cards):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 2.65)

        c = create_card(s6, left, top, Inches(5.7), Inches(2.45), WHITE, BORDER_COLOR)
        tb = s6.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), Inches(5.1), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY if col == 0 else ACCENT_BLUE

        for it in items:
            pi = tf.add_paragraph()
            pi.text = "• " + it
            pi.font.size = Pt(12)
            pi.font.color.rgb = TEXT_DARK
            pi.space_before = Pt(4)

    s6.notes_slide.notes_text_frame.text = (
        "Security is central to StayHive. We implemented both client-side and server-side validation using Joi, "
        "and applied strict middleware boundaries so unauthorized users cannot tamper with listings or reviews."
    )

    # ==========================================
    # SLIDE 7: DATABASE SCHEMA & DATA MODELING
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7, LIGHT_BG)
    add_header(s7, "Database Modeling & Mongoose Schemas")

    schemas = [
        ("Listing Schema",
         [("title", "String (Required)"),
          ("description", "String"),
          ("image", "{ url: String, filename: String }"),
          ("price", "Number (with default & validation)"),
          ("location & country", "String (e.g. Goa, India)"),
          ("geometry", "GeoJSON { type: 'Point', coordinates: [lng, lat] }"),
          ("reviews", "Array of ObjectId refs -> Review model"),
          ("owner", "ObjectId ref -> User model")]),
        ("User Schema",
         [("email", "String (Required, Unique)"),
          ("username", "String (Handled by passport-local-mongoose)"),
          ("hash & salt", "Binary cryptographic hashes"),
          ("createdAt", "Timestamp of registration")]),
        ("Review Schema",
         [("comment", "String (Required)"),
          ("rating", "Number (Min: 1, Max: 5)"),
          ("createdAt", "Date (Defaults to Date.now)"),
          ("author", "ObjectId ref -> User model")])
    ]

    for idx, (name, fields) in enumerate(schemas):
        left = Inches(0.8 + idx * 4.0)
        width = Inches(3.7)
        c = create_card(s7, left, Inches(1.8), width, Inches(5.0), WHITE, BORDER_COLOR)

        # Header box inside card
        hb = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.8), width, Inches(0.6))
        hb.fill.solid()
        hb.fill.fore_color.rgb = PRIMARY if idx == 0 else (ACCENT_BLUE if idx == 1 else DARK_BG)
        hb.line.fill.background()
        
        tb_h = s7.shapes.add_textbox(left, Inches(1.85), width, Inches(0.5))
        p_h = tb_h.text_frame.paragraphs[0]
        p_h.text = name
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = WHITE
        p_h.alignment = PP_ALIGN.CENTER

        # Fields list
        tb = s7.shapes.add_textbox(left + Inches(0.25), Inches(2.55), width - Inches(0.5), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True

        for fn, fd in fields:
            p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
            p.text = f"• {fn}"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = DARK_BG
            
            p_sub = tf.add_paragraph()
            p_sub.text = f"  └ {fd}"
            p_sub.font.size = Pt(11)
            p_sub.font.color.rgb = TEXT_MUTED
            p_sub.space_before = Pt(2)

    s7.notes_slide.notes_text_frame.text = (
        "Here is the database schema design. We utilize relational references in MongoDB using Mongoose ObjectIds. "
        "Notice the GeoJSON point field in the Listing schema, which enables precise spatial plotting on Mapbox."
    )

    # ==========================================
    # SLIDE 8: PRODUCTION DEPLOYMENT & DEVOPS
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8, LIGHT_BG)
    add_header(s8, "Production Deployment & Serverless Architecture")

    deploy_steps = [
        ("⚡ Serverless Express on Vercel",
         "The Express monolith was adapted into an AWS Lambda Serverless Function via @vercel/node. Cold starts and routing are managed seamlessly through vercel.json."),
        ("🌐 MongoDB Atlas Cloud Cluster",
         "Migrated from local MongoDB to an M0 replica set hosted on AWS Mumbai with 0.0.0.0/0 network whitelist and secure SCRAM authentication."),
        ("📦 Static Asset Bundling",
         "Solved serverless asset resolution by bundling views, partials, and styles via includeFiles and routing static assets directly through Express."),
        ("🔐 Environment Variable Isolation",
         "Sensitive API secrets (ATLASDB_URL, CLOUD_NAME, MAP_TOKEN, SECRET) are stored encrypted in Vercel's production environment variables.")
    ]

    for idx, (title, desc) in enumerate(deploy_steps):
        top = Inches(1.8 + idx * 1.25)
        c = create_card(s8, Inches(0.8), top, Inches(11.7), Inches(1.1), WHITE, BORDER_COLOR)
        
        # Circle badge
        badge = s8.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.1), top + Inches(0.25), Inches(0.6), Inches(0.6))
        badge.fill.solid()
        badge.fill.fore_color.rgb = PRIMARY
        badge.line.fill.background()
        
        tb_b = s8.shapes.add_textbox(Inches(1.1), top + Inches(0.28), Inches(0.6), Inches(0.5))
        p_b = tb_b.text_frame.paragraphs[0]
        p_b.text = str(idx + 1)
        p_b.font.size = Pt(16)
        p_b.font.bold = True
        p_b.font.color.rgb = WHITE
        p_b.alignment = PP_ALIGN.CENTER

        tb = s8.shapes.add_textbox(Inches(1.9), top + Inches(0.12), Inches(10.3), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = DARK_BG

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(3)

    s8.notes_slide.notes_text_frame.text = (
        "Deploying an Express app to Vercel requires adapting traditional long-running Node servers into serverless "
        "functions. We configured MongoDB connection caching, bundled EJS templates, and isolated all credentials."
    )

    # ==========================================
    # SLIDE 9: CHALLENGES FACED & SOLUTIONS
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9, LIGHT_BG)
    add_header(s9, "Technical Challenges & Solutions")

    challenges = [
        ("Serverless Asset 404s",
         "Challenge: Vercel serverless functions do not share local filesystem paths by default, causing missing EJS views and 404s on style.css.",
         "Solution: Configured vercel.json with comprehensive includeFiles globs and routed static requests through express.static(__dirname)."),
        ("Database Connection Lifecycle",
         "Challenge: Serverless invocations create multiple concurrent connections, risking MongoDB connection exhaustion.",
         "Solution: Implemented connection state caching (mongoose.connection.readyState) to reuse active connections across serverless requests."),
        ("Mapbox Geocoding Synchronization",
         "Challenge: Addresses with minor misspellings could fail or return empty coordinate arrays, crashing map rendering.",
         "Solution: Added defensive conditional checks on response.body.features and fallback coordinates with null-safe EJS checks."),
        ("Cascade Review Cleanup",
         "Challenge: Deleting a listing left orphaned reviews in the database, cluttering storage.",
         "Solution: Implemented a Mongoose post('findOneAndDelete') hook that automatically triggers Review.deleteMany for all linked IDs.")
    ]

    for idx, (title, chal, sol) in enumerate(challenges):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.8 + row * 2.65)

        c = create_card(s9, left, top, Inches(5.7), Inches(2.45), WHITE, BORDER_COLOR)
        tb = s9.shapes.add_textbox(left + Inches(0.25), top + Inches(0.15), Inches(5.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"⚙️ {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY if col == 0 else ACCENT_BLUE

        pc = tf.add_paragraph()
        pc.text = chal
        pc.font.size = Pt(12)
        pc.font.color.rgb = DARK_BG
        pc.space_before = Pt(6)

        ps = tf.add_paragraph()
        ps.text = sol
        ps.font.size = Pt(12)
        ps.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green
        ps.space_before = Pt(6)

    s9.notes_slide.notes_text_frame.text = (
        "Here are key technical challenges encountered during development and deployment, particularly around "
        "serverless architecture, database pooling, and cascade document cleanup."
    )

    # ==========================================
    # SLIDE 10: FUTURE ROADMAP & CONCLUSION
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10, DARK_BG)

    # Header
    tb_h = s10.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.9))
    tf_h = tb_h.text_frame
    p_cat = tf_h.paragraphs[0]
    p_cat.text = "ROADMAP & CLOSING".upper()
    p_cat.font.size = Pt(12)
    p_cat.font.bold = True
    p_cat.font.color.rgb = PRIMARY

    p_t = tf_h.add_paragraph()
    p_t.text = "Future Roadmap & Project Conclusion"
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE
    p_t.space_before = Pt(4)

    # Left Card: Roadmap
    c_left = create_card(s10, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8), DARK_CARD, None)
    tb_l = s10.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.1), Inches(4.4))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p_rt = tf_l.paragraphs[0]
    p_rt.text = "🚀 Future Enhancements"
    p_rt.font.size = Pt(18)
    p_rt.font.bold = True
    p_rt.font.color.rgb = ACCENT_BLUE

    roadmap_items = [
        "💳 Payment Gateway Integration: Stripe & Razorpay checkout for instant booking payments.",
        "📅 Date Picker & Booking Engine: Interactive availability calendar with reservation confirmation.",
        "💬 Real-Time Messaging: Host-to-guest WebSockets chat for stay coordination.",
        "🔍 Full-Text Elasticsearch: Intelligent search by price range, amenities, and guest count."
    ]
    for rm in roadmap_items:
        p = tf_l.add_paragraph()
        p.text = "• " + rm
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.space_before = Pt(12)

    # Right Card: Conclusion & Live Links
    c_right = create_card(s10, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), DARK_CARD, None)
    tb_r = s10.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.4))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_ct = tf_r.paragraphs[0]
    p_ct.text = "🎉 Project Summary"
    p_ct.font.size = Pt(18)
    p_ct.font.bold = True
    p_ct.font.color.rgb = PRIMARY

    concl_items = [
        "Fully deployed, production-grade web application combining modern cloud services.",
        "Demonstrates complete mastery of the MVC architecture and RESTful design principles.",
        "Enterprise security standards with role-based authorization and salt hashing.",
        "Live on Vercel and accessible worldwide on any device."
    ]
    for ci in concl_items:
        p = tf_r.add_paragraph()
        p.text = "✔ " + ci
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE
        p.space_before = Pt(10)

    # Demo Banner in Right Card
    tb_demo = s10.shapes.add_textbox(Inches(7.1), Inches(5.4), Inches(5.1), Inches(1.2))
    tf_d = tb_demo.text_frame
    p_d = tf_d.paragraphs[0]
    p_d.text = "🌐 Live Demo URL:"
    p_d.font.size = Pt(14)
    p_d.font.bold = True
    p_d.font.color.rgb = PRIMARY

    p_du = tf_d.add_paragraph()
    p_du.text = "https://stay-hive-three.vercel.app/listings"
    p_du.font.size = Pt(13)
    p_du.font.color.rgb = ACCENT_BLUE
    p_du.space_before = Pt(4)

    s10.notes_slide.notes_text_frame.text = (
        "To conclude, StayHive demonstrates an end-to-end full stack development lifecycle, from local MVP "
        "to a live serverless deployment. Thank you for your time, and I would now love to open the floor "
        "to any questions and run a live demonstration of the website!"
    )

    output_path = "StayHive_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
